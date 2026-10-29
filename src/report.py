"""Écriture du rapport : JSON (brut, pour comparer d'une semaine à l'autre), Excel et HTML."""
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERDICTS = ["Active (prouvée)", "Probable", "Douteuse", "Morte"]
VERDICT_COLORS = {"Active (prouvée)": "C6EFCE", "Probable": "E2F0D9", "Douteuse": "FFEB9C", "Morte": "FFC7CE"}


def previous_snapshot(folder: Path, project: str) -> dict | None:
    files = sorted(folder.glob(f"inventaire_{project}_*.json"))
    if not files:
        return None
    try:
        return json.loads(files[-1].read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def diff(prev: dict | None, cur: dict) -> list[dict]:
    """Changements depuis le rapport précédent : règles ajoutées, supprimées, modifiées."""
    if not prev:
        return []
    old = {r["uuid"]: r for r in prev.get("rules", [])}
    new = {r["uuid"]: r for r in cur["rules"]}
    out = []
    for u in new.keys() - old.keys():
        out.append({"type": "Nouvelle règle", "rule": new[u]["name"], "detail": new[u]["state"]})
    for u in old.keys() - new.keys():
        out.append({"type": "Règle disparue", "rule": old[u]["name"], "detail": "supprimée ou hors scope"})
    for u in new.keys() & old.keys():
        a, b = old[u], new[u]
        for key, label in (("name", "Nom"), ("state", "Statut"), ("verdict", "Verdict"), ("updated", "Modifiée le")):
            if a.get(key) != b.get(key):
                out.append({"type": f"{label} changé", "rule": b["name"], "detail": f"{a.get(key)} → {b.get(key)}"})
    return sorted(out, key=lambda c: (c["type"], c["rule"]))


def write_json(data: dict, path: Path) -> None:
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


# --- Excel --------------------------------------------------------------------------

INVENTORY_COLUMNS = [
    ("Nom", "name", 35), ("Statut", "state", 11), ("Verdict", "verdict", 16),
    ("Rôle (ce que fait la règle)", "role", 70), ("Preuve d'utilisation", "evidence", 55),
    ("Dernière action réelle", "last_seen", 13), ("Dernière exécution", "last_run", 13),
    ("Déclenchements (période)", "runs_window", 14), ("Dont avec actions", "acted_window", 12),
    ("Erreurs", "errors", 9), ("Portée", "scope", 13),
    ("Auteur / propriétaire", "author", 22), ("Acteur (exécute en tant que)", "actor", 22),
    ("Éditeurs", "collaborators", 22), ("Créée le", "created", 11), ("Modifiée le", "updated", 11),
    ("Personnes concernées (config.)", "people_config", 35),
    ("Personnes concernées (tickets touchés)", "people_touched", 35),
    ("Tickets exemples", "evidence_issues", 25), ("Remarques", "remarks", 60),
    ("Description", "description", 40), ("Labels", "labels", 15), ("UUID", "uuid", 38),
]


def _cell(value):
    if isinstance(value, list):
        return "\n".join(f"• {v}" for v in value) if value and len(value) > 1 else (value[0] if value else "")
    return value


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def write_excel(data: dict, changes: list[dict], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    rules = data["rules"]

    ws = _sheet(wb, "Inventaire", [c[0] for c in INVENTORY_COLUMNS],
                [[_cell(r.get(c[1])) for c in INVENTORY_COLUMNS] for r in rules],
                [c[2] for c in INVENTORY_COLUMNS])
    verdict_col = [c[1] for c in INVENTORY_COLUMNS].index("verdict") + 1
    for row in ws.iter_rows(min_row=2, min_col=verdict_col, max_col=verdict_col):
        color = VERDICT_COLORS.get(row[0].value)
        if color:
            row[0].fill = PatternFill("solid", fgColor=color)

    _sheet(wb, "Étapes", ["Règle", "Type", "Étape"],
           [[r["name"], s["type"], "    " * s["depth"] + s["text"]] for r in rules for s in r["steps"]],
           [35, 13, 100])
    _sheet(wb, "Remarques", ["Règle", "Statut", "Verdict", "Remarque"],
           [[r["name"], r["state"], r["verdict"], m] for r in rules for m in r["remarks"]],
           [35, 11, 16, 90])
    _sheet(wb, "Changements", ["Type", "Règle", "Détail"],
           [[c["type"], c["rule"], c["detail"]] for c in changes] or [["(premier rapport ou aucun changement)", "", ""]],
           [25, 40, 60])

    counts = Counter(r["verdict"] for r in rules)
    src = data["sources"]
    info = [
        ["Projet", data["project"]],
        ["Généré le", data["generated_at"]],
        ["Fenêtre d'analyse", f"{data['window_days']} jours"],
        ["Source des règles", src["rules"]],
        ["Journal d'audit", f"lu (historique depuis le {src.get('audit_since')})" if src["internal_api"]
         else "non disponible"],
        ["Tickets analysés", src["issues_analyzed"]],
        ["Actions d'automation retrouvées", src["automation_events"]],
        ["Nombre de règles", len(rules)],
        ["Activées", sum(r["state"] == "Activée" for r in rules)],
        *[[f"Verdict : {v}", counts.get(v, 0)] for v in VERDICTS],
        ["Règles avec remarques", sum(bool(r["remarks"]) for r in rules)],
    ]
    _sheet(wb, "Infos", ["Élément", "Valeur"], info, [35, 40])
    wb.move_sheet("Infos", offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    wb.save(path)


# --- HTML ---------------------------------------------------------------------------

HTML_TEMPLATE = """<!doctype html>
<html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Inventaire automations {{ d.project }}</title>
<style>
:root{--bg:#fff;--fg:#1d2433;--muted:#5b6475;--line:#e3e6ec;--card:#f6f7f9;
 --ok:#1f7a3a;--okbg:#dff3e5;--pr:#3d6b1f;--prbg:#eef6e4;--du:#8a5a00;--dubg:#fff1cc;--ko:#a3202a;--kobg:#fde2e4}
@media (prefers-color-scheme:dark){:root{--bg:#14171c;--fg:#e6e8ec;--muted:#9aa3b2;--line:#2a2f38;--card:#1c2027;
 --ok:#7ed39a;--okbg:#173322;--pr:#b5d98f;--prbg:#24301a;--du:#f0c45a;--dubg:#3a2e10;--ko:#f28b93;--kobg:#3b181c}}
body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1200px;margin:0 auto;padding:24px 16px}
h1{font-size:22px;margin:0 0 4px} .muted{color:var(--muted)}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin:18px 0}
.card{background:var(--card);border-radius:8px;padding:10px 12px}.card b{display:block;font-size:22px}
.badge{display:inline-block;padding:1px 8px;border-radius:10px;font-size:12px;font-weight:600;white-space:nowrap}
.v-Active{color:var(--ok);background:var(--okbg)}.v-Probable{color:var(--pr);background:var(--prbg)}
.v-Douteuse{color:var(--du);background:var(--dubg)}.v-Morte{color:var(--ko);background:var(--kobg)}
details.rule{border:1px solid var(--line);border-radius:8px;margin:8px 0;padding:0 12px}
details.rule>summary{cursor:pointer;padding:10px 0;display:flex;gap:10px;align-items:center;flex-wrap:wrap}
details.rule>summary .name{font-weight:600;flex:1;min-width:200px}
dl{display:grid;grid-template-columns:minmax(120px,200px) 1fr;gap:4px 12px;margin:0 0 12px}
dt{color:var(--muted)} dd{margin:0} ul{margin:0;padding-left:18px}
.steps li{list-style:none;margin-left:calc(var(--d)*18px)} .steps{padding-left:0}
.remarks li{color:var(--du)} table{border-collapse:collapse;width:100%} td,th{border-bottom:1px solid var(--line);padding:4px 6px;text-align:left}
@media (max-width:600px){dl{grid-template-columns:1fr}}
</style></head><body><main>
<h1>Inventaire des automations Jira — {{ d.project }}</h1>
<div class="muted">Généré le {{ d.generated_at }} · fenêtre d'analyse {{ d.window_days }} jours ·
 source : {{ d.sources.rules }} · journal d'audit : {{ ("lu depuis le " ~ d.sources.audit_since) if d.sources.internal_api else "non disponible" }}</div>
<div class="cards">
 <div class="card"><b>{{ d.rules|length }}</b>règles</div>
 <div class="card"><b>{{ enabled }}</b>activées</div>
 {% for v in verdicts %}<div class="card"><b>{{ counts.get(v, 0) }}</b>{{ v }}</div>{% endfor %}
 <div class="card"><b>{{ with_remarks }}</b>avec remarques</div>
</div>
{% if changes %}<h2>Changements depuis le dernier rapport</h2>
<table><tr><th>Type</th><th>Règle</th><th>Détail</th></tr>
{% for c in changes %}<tr><td>{{ c.type }}</td><td>{{ c.rule }}</td><td>{{ c.detail }}</td></tr>{% endfor %}</table>{% endif %}
<h2>Règles</h2>
{% for r in d.rules %}<details class="rule"><summary>
 <span class="name">{{ r.name }}</span><span class="badge">{{ r.state }}</span>
 <span class="badge v-{{ r.verdict.split()[0] }}">{{ r.verdict }}</span>
 {% if r.remarks %}<span class="muted">{{ r.remarks|length }} remarque(s)</span>{% endif %}</summary>
 <dl>
  <dt>Rôle</dt><dd>{{ r.role }}</dd>
  <dt>Preuve</dt><dd>{{ r.evidence }}</dd>
  {% if r.runs_window is not none %}<dt>Exécutions ({{ d.window_days }} j)</dt><dd>{{ r.runs_window }} déclenchement(s), dont {{ r.acted_window }} avec actions{% if r.errors %}, {{ r.errors }} en erreur{% endif %} · dernière exécution : {{ r.last_run or "—" }}</dd>{% endif %}
  <dt>Portée</dt><dd>{{ r.scope }}</dd>
  <dt>Auteur / propriétaire</dt><dd>{{ r.author }}</dd>
  <dt>Acteur</dt><dd>{{ r.actor }}</dd>
  {% if r.collaborators %}<dt>Éditeurs</dt><dd>{{ r.collaborators|join(", ") }}</dd>{% endif %}
  <dt>Créée / modifiée</dt><dd>{{ r.created }} / {{ r.updated }}</dd>
  <dt>Personnes citées dans la règle</dt><dd>{{ r.people_config|join(", ") or "—" }}</dd>
  <dt>Personnes des tickets touchés</dt><dd>{{ r.people_touched|join(", ") or "—" }}</dd>
  {% if r.description %}<dt>Description</dt><dd>{{ r.description }}</dd>{% endif %}
  <dt>Étapes</dt><dd><ul class="steps">{% for s in r.steps %}<li style="--d:{{ s.depth }}"><b>{{ s.type }}</b> : {{ s.text }}</li>{% endfor %}</ul></dd>
  {% if r.remarks %}<dt>Remarques</dt><dd><ul class="remarks">{% for m in r.remarks %}<li>{{ m }}</li>{% endfor %}</ul></dd>{% endif %}
 </dl></details>{% endfor %}
</main></body></html>"""


def write_html(data: dict, changes: list[dict], path: Path) -> None:
    env = Environment(autoescape=select_autoescape(default=True, default_for_string=True))
    counts = Counter(r["verdict"] for r in data["rules"])
    html = env.from_string(HTML_TEMPLATE).render(
        d=data, changes=changes, verdicts=VERDICTS, counts=counts,
        enabled=sum(r["state"] == "Activée" for r in data["rules"]),
        with_remarks=sum(bool(r["remarks"]) for r in data["rules"]),
    )
    path.write_text(html, encoding="utf-8")


def write_all(data: dict, out_dir: Path) -> dict[str, Path]:
    folder = out_dir / data["project"]
    folder.mkdir(parents=True, exist_ok=True)
    changes = diff(previous_snapshot(folder, data["project"]), data)
    stamp = datetime.fromisoformat(data["generated_at"]).strftime("%Y-%m-%d_%H%M")
    base = folder / f"inventaire_{data['project']}_{stamp}"
    paths = {"json": base.with_suffix(".json"), "excel": base.with_suffix(".xlsx"), "html": base.with_suffix(".html")}
    write_excel(data, changes, paths["excel"])
    write_html(data, changes, paths["html"])
    write_json(data, paths["json"])  # en dernier : sert de référence pour le prochain diff
    return paths
